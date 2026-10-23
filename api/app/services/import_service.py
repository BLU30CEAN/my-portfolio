from __future__ import annotations

import csv
import io
import uuid
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.entities import AdminEntity, AdminRecord, ImportBatch


def _normalize_key(value: Any) -> str:
    text = str(value or "").strip()
    return text or f"row-{uuid.uuid4().hex[:8]}"


def parse_tabular(file_name: str, raw: bytes) -> tuple[list[str], list[dict[str, str]]]:
    lower = file_name.lower()
    if lower.endswith(".csv"):
        text = raw.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            return [], []
        columns = [c.strip() for c in reader.fieldnames if c]
        rows: list[dict[str, str]] = []
        for row in reader:
            rows.append({k: str(row.get(k) or "").strip() for k in columns})
        return columns, rows

    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        wb = load_workbook(filename=io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        iter_rows = ws.iter_rows(values_only=True)
        header = next(iter_rows, None)
        if not header:
            return [], []
        columns = [str(c or "").strip() or f"col_{i}" for i, c in enumerate(header)]
        rows = []
        for line in iter_rows:
            if line is None:
                continue
            if all(v is None or str(v).strip() == "" for v in line):
                continue
            row = {
                columns[i]: str(line[i] if i < len(line) else "").strip()
                for i in range(len(columns))
            }
            rows.append(row)
        return columns, rows

    raise ValueError("지원 형식: .csv, .xlsx")


def infer_column_schema(columns: list[str]) -> list[dict[str, str]]:
    return [{"key": c, "label": c, "kind": "text"} for c in columns]


async def apply_import(
    session: AsyncSession,
    entity: AdminEntity,
    *,
    file_name: str,
    mode: str,
    raw: bytes,
    key_column: str | None,
) -> ImportBatch:
    columns, rows = parse_tabular(file_name, raw)
    if not columns:
        raise ValueError("헤더가 없거나 데이터가 비어 있습니다.")

    if not entity.column_schema:
        entity.column_schema = infer_column_schema(columns)

    key_col = key_column or columns[0]
    inserted = updated = skipped = 0
    errors: list[str] = []

    if mode == "replace":
        await session.execute(
            delete(AdminRecord).where(AdminRecord.entity_id == entity.id)
        )

    existing: dict[str, AdminRecord] = {}
    if mode in ("upsert", "insert"):
        result = await session.execute(
            select(AdminRecord).where(AdminRecord.entity_id == entity.id)
        )
        for rec in result.scalars().all():
            existing[rec.external_key] = rec

    known_keys = set(existing.keys())

    for idx, row in enumerate(rows, start=2):
        try:
            ext = _normalize_key(row.get(key_col))
            if mode == "insert" and ext in known_keys:
                skipped += 1
                continue

            if mode == "upsert" and ext in existing:
                existing[ext].data = row
                updated += 1
                continue

            session.add(
                AdminRecord(entity_id=entity.id, external_key=ext, data=row)
            )
            inserted += 1
            known_keys.add(ext)
        except Exception as exc:  # noqa: BLE001 — row-level import audit
            errors.append(f"행 {idx}: {exc}")

    batch = ImportBatch(
        entity_id=entity.id,
        file_name=file_name,
        mode=mode,
        row_count=len(rows),
        inserted=inserted,
        updated=updated,
        skipped=skipped,
        errors=errors[:50],
    )
    session.add(batch)
    await session.commit()
    await session.refresh(batch)
    return batch
